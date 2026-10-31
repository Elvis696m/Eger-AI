import os
import re
import hmac
import threading

from io import BytesIO
from pathlib import Path
from functools import wraps

from flask import Flask, request, jsonify
from flask_sqlalchemy import SQLAlchemy
from flask_cors import CORS
from werkzeug.utils import secure_filename

from pypdf import PdfReader
from docx import Document as WordDocument

from sklearn.feature_extraction.text import (
    TfidfVectorizer,
    ENGLISH_STOP_WORDS
)
from sklearn.metrics.pairwise import cosine_similarity


# ==================================================
# 1. FLASK CONFIGURATION
# ==================================================

app = Flask(__name__)
CORS(app)

app.config["MAX_CONTENT_LENGTH"] = 10 * 1024 * 1024

BASE_DIR = os.path.abspath(os.path.dirname(__file__))
INSTANCE_DIR = os.path.join(BASE_DIR, "instance")
os.makedirs(INSTANCE_DIR, exist_ok=True)

DATABASE_PATH = os.path.join(INSTANCE_DIR, "egerai.db")

app.config["SQLALCHEMY_DATABASE_URI"] = (
    "sqlite:///" + DATABASE_PATH.replace("\\", "/")
)
app.config["SQLALCHEMY_TRACK_MODIFICATIONS"] = False

db = SQLAlchemy(app)


# ==================================================
# 2. SEARCH SETTINGS
# ==================================================

# Words that describe the *request*, not the topic.
QUESTION_FILLERS = {
    "define", "explain", "describe", "tell", "give", "show", "list",
    "please", "briefly", "meaning", "discuss", "outline", "state",
    "elaborate", "mean", "means", "about", "bro", "bruh", "dude"
}

STOP_WORDS = list(ENGLISH_STOP_WORDS | QUESTION_FILLERS)

# Abbreviations students use. Add your course abbreviations here.
ABBREVIATIONS = {
    "plf": "plf precision livestock farming",
}

# Below this score: no answer. Between the two: answer with low confidence.
MINIMUM_SCORE = 0.05
CONFIDENT_SCORE = 0.12
SUGGESTION_FLOOR = 0.02


# ==================================================
# 3. ADMIN AUTHENTICATION
# ==================================================

def admin_required(route_function):

    @wraps(route_function)
    def wrapper(*args, **kwargs):

        if request.method == "OPTIONS":
            return route_function(*args, **kwargs)

        admin_password = os.environ.get("ADMIN_PASSWORD")

        if not admin_password:
            app.logger.error("ADMIN_PASSWORD is not configured.")
            return jsonify({
                "error": "Administrator authentication is not configured."
            }), 503

        authorization = request.headers.get("Authorization", "")
        supplied_password = ""

        if authorization.startswith("Bearer "):
            supplied_password = authorization[len("Bearer "):]

        if not hmac.compare_digest(supplied_password, admin_password):
            return jsonify({"error": "Unauthorized."}), 401

        return route_function(*args, **kwargs)

    return wrapper


# ==================================================
# 4. DATABASE MODELS
# ==================================================

class Knowledge(db.Model):

    __tablename__ = "knowledge"

    id = db.Column(db.Integer, primary_key=True)
    question = db.Column(db.String(200), nullable=False)
    keywords = db.Column(db.String(500), nullable=False)
    answer = db.Column(db.Text, nullable=False)


class Document(db.Model):

    __tablename__ = "documents"

    id = db.Column(db.Integer, primary_key=True)
    title = db.Column(db.String(200), nullable=False)
    category = db.Column(db.String(100), nullable=False)
    filename = db.Column(db.String(255), nullable=False)
    extracted_text = db.Column(db.Text, nullable=False)
    uploaded_at = db.Column(
        db.DateTime,
        server_default=db.func.current_timestamp()
    )


# ==================================================
# 5. INITIALIZE DATABASE
# ==================================================

with app.app_context():

    db.create_all()

    print("====================================")
    print("EGER AI DATABASE INITIALIZED")
    print("====================================")
    print("Knowledge records:", Knowledge.query.count())
    print("Stored documents:", Document.query.count())


# ==================================================
# 6. CLEAN TEXT
# ==================================================

def clean_text(text):

    text = str(text or "").lower()
    text = re.sub(r"[^a-z0-9\s]", " ", text)
    text = re.sub(r"\s+", " ", text)

    return text.strip()


def expand_abbreviations(text):

    words = clean_text(text).split()

    return " ".join(
        ABBREVIATIONS.get(word, word)
        for word in words
    )


# ==================================================
# 7. REMOVE CONVERSATIONAL FILLERS
# ==================================================

def normalize_question(question):

    question = str(question or "").strip()

    question = re.sub(
        r"^(bro+|bruh+|dude|mate|man)[\s,!?]*",
        "",
        question,
        flags=re.IGNORECASE
    )

    question = re.sub(
        r"[\s,!?]+\b(bro+|bruh+|dude|mate|man)\b[.!?]*$",
        "",
        question,
        flags=re.IGNORECASE
    )

    return question.strip()


# ==================================================
# 8. RECOGNIZE FOLLOW-UP INTENT
# ==================================================

def recognize_followup(question):

    cleaned = clean_text(question)

    followup_patterns = {

        "simplify": [
            "simplify", "simplify that", "simplify it",
            "make it simple", "make it easier", "explain simply",
            "in simple english", "easy english"
        ],

        "examples": [
            "examples", "give examples", "some examples",
            "show examples", "example please", "give me examples"
        ],

        "explain": [
            "explain that", "explain it", "explain further",
            "elaborate", "elaborate on that", "tell me more",
            "more details", "go deeper", "expand on that"
        ],

        "why": [
            "why", "why is that", "why is that so", "how come"
        ],

        "continue": [
            "continue", "go on", "proceed", "continue explaining"
        ],

        "summarize": [
            "summarize", "summarise", "short summary",
            "summarize that", "summarise that"
        ],

        "clarify": [
            "what do you mean", "what do you mean by that",
            "what does that mean", "clarify",
            "i dont understand", "i do not understand"
        ]
    }

    for intent, phrases in followup_patterns.items():
        if cleaned in phrases:
            return intent

    return None


# ==================================================
# 9. BUILD CONTEXTUAL QUESTION
# ==================================================

def build_contextual_question(
    user_question,
    previous_question,
    previous_answer
):

    intent = recognize_followup(user_question)

    if not intent or not previous_question:
        return user_question, None

    intent_keywords = {
        "simplify": "simple explanation",
        "examples": "examples applications",
        "explain": "explanation details",
        "why": "reasons causes importance",
        "continue": "further explanation",
        "summarize": "summary main points",
        "clarify": "definition explanation"
    }

    contextual_question = (
        previous_question + " " + intent_keywords.get(intent, "")
    )

    return contextual_question, intent


# ==================================================
# 10. SPLIT DOCUMENTS INTO PASSAGES
# Splits on slide/page breaks (blank lines) first, then
# groups whole lines up to chunk_size words, so passages
# do not cross slide boundaries.
# ==================================================

def split_into_passages(
    text,
    chunk_size=110,
    overlap=30
):

    if chunk_size <= 0 or overlap < 0 or overlap >= chunk_size:
        raise ValueError("Invalid passage chunk size or overlap.")

    blocks = [
        block for block in re.split(r"\n\s*\n", text or "")
        if block.strip()
    ]

    passages = []

    for block in blocks:

        # Break the block into line-sized units.
        units = []

        for line in block.splitlines():

            words = line.split()

            if not words:
                continue

            for i in range(0, len(words), chunk_size):
                units.append(" ".join(words[i:i + chunk_size]))

        current = []
        count = 0

        for unit in units:

            unit_count = len(unit.split())

            if current and count + unit_count > chunk_size:

                passages.append("\n".join(current))

                # Carry the last few lines over as overlap.
                tail = []
                tail_count = 0

                for previous in reversed(current):

                    previous_count = len(previous.split())

                    if tail_count + previous_count > overlap:
                        break

                    tail.insert(0, previous)
                    tail_count += previous_count

                current = tail
                count = tail_count

            current.append(unit)
            count += unit_count

        if current:
            passages.append("\n".join(current))

    return passages


# ==================================================
# 11. CACHED DOCUMENT INDEX
# The index is built from passages only (the question is
# never part of the fit) and rebuilt only when documents
# change.
# ==================================================

_index = {
    "key": None,
    "vectorizer": None,
    "matrix": None,
    "passages": []
}

_index_lock = threading.Lock()


def invalidate_index():

    with _index_lock:
        _index["key"] = None


def build_index():

    with _index_lock:

        documents = Document.query.order_by(Document.id).all()

        key = (
            len(documents),
            max((d.id for d in documents), default=0)
        )

        if _index["key"] == key:
            return _index

        passages = []

        for document in documents:

            for number, passage in enumerate(
                split_into_passages(document.extracted_text)
            ):

                cleaned = clean_text(passage)

                if cleaned:
                    passages.append({
                        "document_id": document.id,
                        "title": document.title,
                        "category": document.category,
                        "passage_number": number + 1,
                        "text": passage,
                        "cleaned_text": cleaned
                    })

        vectorizer = None
        matrix = None

        if passages:
            try:
                vectorizer = TfidfVectorizer(
                    stop_words=STOP_WORDS,
                    ngram_range=(1, 2),
                    sublinear_tf=True
                )
                matrix = vectorizer.fit_transform(
                    [p["cleaned_text"] for p in passages]
                )
            except ValueError:
                vectorizer = None
                matrix = None

        _index.update(
            key=key,
            vectorizer=vectorizer,
            matrix=matrix,
            passages=passages
        )

        return _index


# ==================================================
# 12. SEARCH UNIVERSITY DOCUMENTS
# Returns {"best": match or None, "suggestions": [...]}
# ==================================================

def search_documents(user_question):

    empty = {"best": None, "suggestions": []}

    index = build_index()

    if not index["passages"] or index["vectorizer"] is None:
        return empty

    question = expand_abbreviations(user_question)

    if not question:
        return empty

    question_vector = index["vectorizer"].transform([question])

    # None of the question's words exist in any document.
    if question_vector.nnz == 0:
        return empty

    similarities = cosine_similarity(
        question_vector,
        index["matrix"]
    ).flatten()

    ranked = similarities.argsort()[::-1]

    best_index = int(ranked[0])
    best_score = float(similarities[best_index])

    if best_score < MINIMUM_SCORE:

        suggestions = []

        for i in ranked[:3]:

            score = float(similarities[int(i)])

            if score < SUGGESTION_FLOOR:
                break

            item = index["passages"][int(i)]

            suggestions.append({
                "document_title": item["title"],
                "passage_number": item["passage_number"],
                "similarity": round(score, 4),
                "preview": item["text"][:300]
            })

        return {"best": None, "suggestions": suggestions}

    best = index["passages"][best_index].copy()
    best["similarity"] = round(best_score, 4)
    best["confidence"] = (
        "high" if best_score >= CONFIDENT_SCORE else "low"
    )

    # Attach the next passage from the same document so lists
    # that run past the chunk boundary are not cut off.
    best["next_text"] = None

    next_index = best_index + 1

    if next_index < len(index["passages"]):

        following = index["passages"][next_index]

        if following["document_id"] == best["document_id"]:
            best["next_text"] = following["text"]

    return {"best": best, "suggestions": []}


# ==================================================
# 13. SEARCH KNOWLEDGE DATABASE
# ==================================================

def search_knowledge(user_question):

    cleaned_question = clean_text(user_question)
    question_words = set(cleaned_question.split())

    if not question_words:
        return None

    best_answer = None
    highest_score = 0

    for item in Knowledge.query.all():

        stored_question = clean_text(item.question)
        stored_keywords = clean_text(item.keywords)

        question_matches = len(
            question_words & set(stored_question.split())
        )

        keyword_matches = len(
            question_words & set(stored_keywords.split())
        )

        score = question_matches * 2 + keyword_matches * 3

        if stored_question and stored_question in cleaned_question:
            score += 5

        if score > highest_score:
            highest_score = score
            best_answer = item.answer

    if best_answer and highest_score >= 5:
        return {"answer": best_answer, "score": highest_score}

    return None


# ==================================================
# 14. HOME ROUTE
# ==================================================

@app.route("/", methods=["GET"])
def home():

    return jsonify({
        "message": "Eger AI Flask server is working!",
        "status": "online",
        "document_search": "TF-IDF + cosine similarity (cached index)",
        "conversation_memory": "enabled"
    })


# ==================================================
# 15. ASK EGER AI
# ==================================================

@app.route("/ask", methods=["POST", "OPTIONS"])
def ask():

    if request.method == "OPTIONS":
        return "", 204

    data = request.get_json(silent=True) or {}

    user_question = data.get("question", "")
    previous_question = data.get("previousQuestion", "")
    previous_answer = data.get("previousAnswer", "")

    # A. VALIDATE INPUT

    if not isinstance(user_question, str):
        return jsonify({"error": "Question must be text."}), 400

    if not isinstance(previous_question, str):
        previous_question = ""

    if not isinstance(previous_answer, str):
        previous_answer = ""

    user_question = user_question.strip()

    if not user_question:
        return jsonify({"error": "Please provide a question."}), 400

    # B. NORMALIZE QUESTION

    normalized_question = normalize_question(user_question)

    if not normalized_question:
        normalized_question = user_question

    # C. RECOGNIZE FOLLOW-UP

    contextual_question, followup_intent = build_contextual_question(
        normalized_question,
        previous_question,
        previous_answer
    )

    is_followup = followup_intent is not None

    # D. SEARCH DOCUMENTS

    search = search_documents(contextual_question)
    best = search["best"]

    app.logger.info(
        "Question: %s | Follow-up: %s | Best score: %s",
        normalized_question,
        followup_intent,
        best["similarity"] if best else None
    )

    if best:

        answer = best["text"]

        if best["next_text"]:
            answer = answer + "\n\n" + best["next_text"]

        return jsonify({
            "answer": answer,
            "status": "success",
            "source": "document",
            "document_id": best["document_id"],
            "document_title": best["title"],
            "category": best["category"],
            "passage_number": best["passage_number"],
            "similarity": best["similarity"],
            "confidence": best["confidence"],
            "followup": is_followup,
            "intent": followup_intent
        })

    # E. SEARCH KNOWLEDGE DATABASE

    knowledge_result = search_knowledge(contextual_question)

    if knowledge_result:

        return jsonify({
            "answer": knowledge_result["answer"],
            "status": "success",
            "source": "knowledge",
            "followup": is_followup,
            "intent": followup_intent
        })

    # F. PREVIOUS ANSWER FALLBACK

    if is_followup and previous_answer:

        fallback_messages = {
            "simplify": (
                "I don't have another passage to simplify "
                "right now. Here is the explanation I found:"
            ),
            "examples": (
                "I couldn't find a separate examples section. "
                "Here is the information available:"
            ),
            "explain": (
                "I couldn't find additional information "
                "in the uploaded documents. Here is the "
                "previous explanation:"
            ),
            "why": (
                "I couldn't find further information "
                "explaining why. Here is the previous answer:"
            ),
            "continue": (
                "I couldn't find another relevant passage. "
                "Here is the information available:"
            ),
            "summarize": (
                "Here is the information available to summarize:"
            ),
            "clarify": (
                "Let me refer back to the previous answer:"
            )
        }

        return jsonify({
            "answer": (
                fallback_messages.get(
                    followup_intent,
                    "Here is the previous answer:"
                )
                + "\n\n"
                + previous_answer
            ),
            "status": "success",
            "source": "conversation_memory",
            "followup": True,
            "intent": followup_intent
        })

    # G. NO ANSWER FOUND (with closest passages, if any)

    return jsonify({
        "answer": (
            "I'm sorry, I couldn't find a sufficiently "
            "relevant answer in the university documents "
            "or knowledge base. Please try different "
            "wording or contact the relevant Egerton "
            "University office."
        ),
        "status": "not_found",
        "source": None,
        "followup": False,
        "suggestions": search["suggestions"]
    })


# ==================================================
# 16. ADD KNOWLEDGE (ADMIN PASSWORD REQUIRED)
# ==================================================

@app.route("/add", methods=["POST", "OPTIONS"])
@admin_required
def add_knowledge():

    if request.method == "OPTIONS":
        return "", 204

    data = request.get_json(silent=True) or {}

    question = data.get("question", "")
    keywords = data.get("keywords", "")
    answer = data.get("answer", "")

    if not all(
        isinstance(value, str)
        for value in (question, keywords, answer)
    ):
        return jsonify({"error": "All fields must contain text."}), 400

    question = question.strip()
    keywords = keywords.strip()
    answer = answer.strip()

    if not question or not keywords or not answer:
        return jsonify({
            "error": "Question, keywords, and answer are all required."
        }), 400

    if len(question) > 200:
        return jsonify({
            "error": "Question must not exceed 200 characters."
        }), 400

    if len(keywords) > 500:
        return jsonify({
            "error": "Keywords must not exceed 500 characters."
        }), 400

    try:

        new_knowledge = Knowledge(
            question=question,
            keywords=keywords,
            answer=answer
        )

        db.session.add(new_knowledge)
        db.session.commit()

        return jsonify({
            "message": "Knowledge added successfully!",
            "id": new_knowledge.id
        }), 201

    except Exception:

        db.session.rollback()
        app.logger.exception("Failed to save knowledge.")

        return jsonify({"error": "Failed to save knowledge."}), 500


# ==================================================
# 17. TEXT EXTRACTION
# Pages are joined with a blank line so the passage
# splitter can respect slide/page boundaries.
# ==================================================

def extract_pdf_text(file_bytes):

    pdf = PdfReader(BytesIO(file_bytes))

    if pdf.is_encrypted:
        raise ValueError("This PDF is password-protected.")

    extracted_pages = []

    for page in pdf.pages:

        page_text = page.extract_text() or ""

        if page_text.strip():
            extracted_pages.append(page_text.strip())

    return "\n\n".join(extracted_pages)


def extract_docx_text(file_bytes):

    document = WordDocument(BytesIO(file_bytes))

    extracted_parts = []

    for paragraph in document.paragraphs:

        if paragraph.text.strip():
            extracted_parts.append(paragraph.text)

    for table in document.tables:

        for row in table.rows:

            row_text = " | ".join(
                cell.text.strip() for cell in row.cells
            )

            if row_text.strip():
                extracted_parts.append(row_text)

    return "\n".join(extracted_parts)


def extract_txt_text(file_bytes):

    return file_bytes.decode("utf-8-sig")


# ==================================================
# 18. UPLOAD UNIVERSITY DOCUMENT (ADMIN PASSWORD REQUIRED)
# ==================================================

@app.route("/upload", methods=["POST", "OPTIONS"])
@admin_required
def upload_document():

    if request.method == "OPTIONS":
        return "", 204

    title = request.form.get("title", "").strip()
    category = request.form.get("category", "").strip()
    uploaded_file = request.files.get("document")

    if not title or not category or not uploaded_file:
        return jsonify({
            "error": "Title, category, and document are all required."
        }), 400

    if len(title) > 200:
        return jsonify({
            "error": "Title must not exceed 200 characters."
        }), 400

    if len(category) > 100:
        return jsonify({
            "error": "Category must not exceed 100 characters."
        }), 400

    original_filename = secure_filename(uploaded_file.filename or "")

    if not original_filename:
        return jsonify({"error": "Please select a valid filename."}), 400

    extension = Path(original_filename).suffix.lower()

    if extension not in {".pdf", ".docx", ".txt"}:
        return jsonify({
            "error": "Only PDF, DOCX, and TXT files are supported."
        }), 400

    try:

        file_bytes = uploaded_file.read()

        if not file_bytes:
            return jsonify({"error": "The uploaded file is empty."}), 400

        if extension == ".pdf":
            extracted_text = extract_pdf_text(file_bytes)
        elif extension == ".docx":
            extracted_text = extract_docx_text(file_bytes)
        else:
            extracted_text = extract_txt_text(file_bytes)

        extracted_text = extracted_text.strip()

        if not extracted_text:
            return jsonify({
                "error": (
                    "No readable text was found. "
                    "Scanned PDFs may require OCR."
                )
            }), 400

        new_document = Document(
            title=title,
            category=category,
            filename=original_filename,
            extracted_text=extracted_text
        )

        db.session.add(new_document)
        db.session.commit()

        # Make the new document searchable immediately.
        invalidate_index()

        return jsonify({
            "message": "Document text extracted and saved successfully!",
            "document_id": new_document.id,
            "title": new_document.title,
            "category": new_document.category,
            "characters_extracted": len(extracted_text)
        }), 201

    except (ValueError, UnicodeDecodeError) as error:

        db.session.rollback()

        return jsonify({"error": str(error)}), 400

    except Exception:

        db.session.rollback()
        app.logger.exception("Document upload failed.")

        return jsonify({
            "error": (
                "The document could not be processed. "
                "Check the file and try again."
            )
        }), 500


# ==================================================
# 19. FILE SIZE ERROR
# ==================================================

@app.errorhandler(413)
def file_too_large(error):

    return jsonify({
        "error": "File too large. Maximum size is 10 MB."
    }), 413


# ==================================================
# 20. RUN THE APPLICATION
# ==================================================

if __name__ == "__main__":

    app.run(
        host="127.0.0.1",
        port=5000,
        debug=True
    )